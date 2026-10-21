from openpyxl.workbook.defined_name import DefinedName
# ---- round 2: goal-linked inputs, debt that takes longest, C10 / C13 app method ----
PIN="'Plan inputs'!$N$111:$N$132"
def goalcol(col,key): return f"IFERROR(INDEX('Plan inputs'!${col}$111:${col}$132,MATCH(\"{key}\",{PIN},0)),0)"
def relink(ws_,r_,has,val,label,tag):
    g=ws_.cell(r_,7); orig=g.value; fb=f'IF(Fill_Example="Yes",F{r_},"")'
    assert isinstance(orig,str) and orig.count(fb)==1,(ws_.title,r_,orig)
    g.value=orig.replace(fb,f'IF({has},{val},{fb})'); linked.append((ws_.title[:3],r_,label,tag))
HS=goalcol('R','home'); HA=goalcol('Q','home'); RA=goalcol('Q','retire')
relink(wb['C01 Borrowing'],8,f'{HS}>0',HS,'Deposit saved','Plan goals: home goal, saved so far')
relink(wb['C03 Deposit'],9,f'{HS}>0',HS,'Saved so far','Plan goals: home goal, saved so far')
relink(wb['C07 Rent vs buy'],8,f'{HA}>0',HA,'Deposit','Plan goals: home goal, amount')
relink(wb['C12 Retirement projection'],25,f'{RA}>0',RA,'Desired yearly income in retirement','Plan goals: retire goal, yearly income')
relink(wb['C12 Retirement projection'],26,'OR(Has_sp=1,Has_spYears=1,Has_a_spWeek=1)','ROUND(PR_SPweek*PR_Weeks*F_SpF,0)','Other yearly income from 66 (State Pension)','Plan pay & tax: State Pension share x full yearly rate')
# C27: the debt (card or loan) that takes longest to clear
CM='IFERROR(NPER(D_iC,-F_CardPayM,F_CardBal),1E9)'; LM='IFERROR(NPER(D_iL,-F_LoanPayM,F_LoanBal),1E9)'
c27=wb['C27 Debt repayment']
relink(c27,6,'F_Debt>0',f'IF(F_CardBal>0,IF(F_LoanBal>0,IF({CM}>={LM},F_CardBal,F_LoanBal),F_CardBal),F_LoanBal)','Balance','Plan inputs: the debt that takes longest to clear')
relink(c27,8,'F_Debt>0',f'IF(F_CardBal>0,IF(F_LoanBal>0,IF({CM}>={LM},F_CardPayM,F_LoanPayM),F_CardPayM),F_LoanPayM)','Monthly payment',"Plan inputs: that debt's repayment")
relink(c27,7,'AND(F_Debt>0,IF(F_CardBal>0,IF(F_LoanBal>0,IF(%s>=%s,In_cardRate>0,In_loanRate>0),In_cardRate>0),In_loanRate>0))'%(CM,LM),f'IF(F_CardBal>0,IF(F_LoanBal>0,IF({CM}>={LM},In_cardRate,In_loanRate),In_cardRate),In_loanRate)','Interest (APR)','Plan inputs: that debt\'s rate, when you gave it')
c27['E6']=str(c27['E6'].value)  # keep note
# C10: waiting money earns the cash rate (the app's method)
c10=wb['C10 Lump sum growth']
c10['C33']='=IF(Need_3<>"",Need_3,Used_Amount*(1+AS_Cash)^MIN(Wait_Years,Years)*(1+Growth_After_Fees)^MAX(0,Years-Wait_Years))'
c10['B42']='• Value = amount × (1 + growth after fees)^years. If you start after waiting, the money earns the cash savings rate while it waits (before DIRT), then grows at the fund rate for the years left.'
c10['B45']='• With the wait as long as or longer than the years, the money only earns the cash rate.'
c10['B63']='Note: while it waits, the money earns the cash savings rate (your Plan assumption "Cash savings earn"), as in the app. If growth after fees is no more than the cash rate, waiting costs nothing or less.'
# C13: contributions rise with pay (the app's method)
c13=wb['C13 Contribution impact']
from copy import copy as _cp
for col in range(2,8):
    src=c13.cell(22,col); dst=c13.cell(23,col); dst.font=_cp(src.font); dst.fill=_cp(src.fill); dst.border=_cp(src.border); dst.alignment=_cp(src.alignment); dst.number_format=src.number_format
c13['B23']='Pay rises a year (contributions rise with pay)'; c13['D23']='%'; c13['E23']='Type 3 · LifeMap standard · LifeMap uses 3% a year (Irish pay has grown about 3%-4% a year recently; LifeMap Standard set). Change it if you want to use your own. Allowed: 0% to 6% · steps of 0.25%'
c13['F23']='=AS_Wage'; c13['G23']='=IF(ISBLANK(C23),F23,C23)'; c13['C23'].fill=_cp(c13['C22'].fill); c13['C23'].number_format='0.00%'; c13['F23'].number_format='0.00%'; c13['G23'].number_format='0.00%'
c13.defined_names['Wage']=DefinedName('Wage',attr_text="'C13 Contribution impact'!$G$23")
c13['C29']='=IF(Need_1<>"",Need_1,IF(ABS(Used_Growth_Rate-Wage)<1E-12,12*Extra_Per_Month*Years*(1+Used_Growth_Rate)^(Years-1),12*Extra_Per_Month*((1+Used_Growth_Rate)^Years-(1+Wage)^Years)/(Used_Growth_Rate-Wage)))'
c13['B39']='• Extra a month = salary × increase ÷ 12 in the first year. The extra rises with pay each year (the pay rise on this sheet, from your Plan assumption "Pay rises"), so Value = 12 × extra × ((1+growth)^years − (1+pay rise)^years) ÷ (growth − pay rise).'
c13['B58']='Note: the extra contribution rises with pay each year, as in the app and your plan. Tax relief shown is for the first year.'
c13['C32'].value=c13['C32'].value
json.dump(linked,open(SC+'linked.json','w'))
