"""Adds the stress test to the v3.0 workbook: a side run of the plan (copies of the four engine sheets with only the
stress-year growth changed) and the 'Plan stress test' sheet. Usage: stress_v3.py in.xlsx out.xlsx"""
import sys, re, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.hyperlink import Hyperlink
src, out = sys.argv[1:3]
wb = openpyxl.load_workbook(src)
CL = {'Plan pay & tax': 'Stress pay & tax', 'Plan holdings': 'Stress holdings', 'Plan goals': 'Stress goals', 'Plan cashflow': 'Stress cashflow'}
# names that point into the four engine sheets: inside the copies they must point at the copies
NM = {}
for n, d in wb.defined_names.items():
    m = re.match(r"^'?([^'!]+)'?!(.+)$", d.attr_text or '')
    if m and m.group(1) in CL: NM[n] = "'%s'!%s" % (CL[m.group(1)], m.group(2))
name_re = re.compile(r"(?<![A-Za-z0-9_.!'$])(" + '|'.join(sorted(map(re.escape, NM), key=len, reverse=True)) + r")(?![A-Za-z0-9_.(!])")
sheet_re = re.compile("|".join("'%s'!" % re.escape(k) for k in CL))
def redirect(f):
    parts = f.split('"')           # odd parts are text constants: leave them
    for i in range(0, len(parts), 2):
        p = sheet_re.sub(lambda m: "'%s'!" % CL[m.group(0)[1:-2]], parts[i])
        parts[i] = name_re.sub(lambda m: NM[m.group(1)], p)
    return '"'.join(parts)
cnt = {'pen': 0, 'inv': 0, 'pot': 0}; slots = set()
def stress_edit(sheet, cell):
    v = cell.value
    if sheet == 'Plan pay & tax' and cell.column_letter == 'AD' and 234 <= cell.row <= 323:
        n = v.count('(1+F_PenG)') + v.count('(1+AS_PenRet)'); assert n == 3, (cell.coordinate, v)
        v = v.replace('(1+F_PenG)', 'IF($B%d=ST_T,1-ST_Fall,1+F_PenG)' % cell.row).replace('(1+AS_PenRet)', 'IF($B%d=ST_T,1-ST_Fall,1+AS_PenRet)' % cell.row); cnt['pen'] += 1
    elif sheet == 'Plan holdings' and cell.column_letter == 'BQ' and cell.row >= 7:
        assert v == '=IF(H_TaxEvt=1,H_InvG,AS_Inv)', (cell.coordinate, v)
        v = '=IF($B%d=ST_T,-ST_Fall,IF(H_TaxEvt=1,H_InvG,AS_Inv))' % cell.row; cnt['inv'] += 1
    elif sheet == 'Plan cashflow' and cell.row >= 8 and isinstance(v, str) and re.search(r"\*\(1\+Slot\d+_Rate\)(?!\^)", v) and cell.column_letter not in ('QM',) or (sheet == 'Plan cashflow' and cell.row >= 8 and isinstance(v, str) and re.search(r"\*\(1\+Slot\d+_Rate\),0\)$", v)):
        k = re.search(r"\*\(1\+Slot(\d+)_Rate\)(?!\^)", v)
        if k:
            n = k.group(1); v = v[:k.start()] + "*IF(AND($B%d=ST_T,Slot%s_Rate=AS_Inv,AS_Inv<>AS_Cash),1-ST_Fall,1+Slot%s_Rate)" % (cell.row, n, n) + v[k.end():]; cnt['pot'] += 1; slots.add(int(n))
    return v
for sh, new in CL.items():
    ws = wb[sh]; c = wb.copy_worksheet(ws); c.title = new
    for row in c.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith('='):
                v = stress_edit(sh, cell) if True else v
                cell.value = redirect(v)
    c['B1'] = new + ': a copy of ' + sh + " for the stress test; only the growth in the stress year differs. Do not edit."
    c.sheet_state = 'hidden'; c.sheet_properties.tabColor = '999999'
print('edited', cnt, sorted(slots))
# ---- the Plan stress test sheet
st = wb.create_sheet('Plan stress test')
H = Font(bold=True, size=14); B = Font(bold=True)
st.column_dimensions['A'].width = 4; st.column_dimensions['B'].width = 52; st.column_dimensions['C'].width = 18; st.column_dimensions['D'].width = 80
st['A1'] = '≡'; st['A1'].hyperlink = Hyperlink(ref='A1', location="'READ ME'!A1", display='≡'); st['A1'].font = Font(bold=True, color='0563C1')
st['B1'] = 'Plan stress test: if markets fell'; st['B1'].font = H
st['B2'] = ('A side run of your plan (the plan itself does not change). In the year before the retirement age, pension pots and invested money fall by the share below '
            'instead of growing. Cash, savings accounts, the emergency fund and assets that move with prices do not fall. Everything else follows the plan.')
st['B2'].alignment = Alignment(wrap_text=True, vertical='top'); st.merge_cells('B2:D2'); st.row_dimensions[2].height = 48
rows = [
 ('ST_Fall', 'How much markets fall (blue cell: type another share)', 0.2, 'Same as STRESS_FALL in the app (20%).'),
 ('ST_T', 'Stress year: years from the first plan year', '=MAX(0,F_R-1-F_Age)', 'The year before the retirement age, or this year if you are already retired.'),
 ('ST_Age', 'Your age in the stress year', '=F_Age+ST_T', ''),
 ('ST_Year', 'Calendar year of the stress year', "=INDEX('Plan holdings'!$E$8:$E$97,ST_T+1)", ''),
 ('ST_RunOutAge', 'Age your money runs short, after the fall', '=IF(COUNT($F$4:$F$93)=0,"",MIN($F$4:$F$93))', 'First year living costs and repayments are short by more than max(€500, 2% of that year\'s needs). Blank = lasts to the end of the plan.'),
 ('ST_RunOutAgeBase', 'Age your money runs short, in your plan', '=IF(COUNT($G$4:$G$93)=0,"",MIN($G$4:$G$93))', ''),
 ('ST_PotAtRet', 'Pension pot at retirement, after the fall (€)', "=ROUND(INDEX('Stress pay & tax'!$AD$234:$AD$323,MIN(ST_T,89)+1),0)", ''),
 ('ST_PotAtRetBase', 'Pension pot at retirement, in your plan (€)', "=ROUND(INDEX('Plan pay & tax'!$AD$234:$AD$323,MIN(ST_T,89)+1),0)", ''),
]
r0 = 4
st.cell(r0 - 1, 2, 'RESULTS').font = B; st.cell(r0 - 1, 3, 'Value').font = B; st.cell(r0 - 1, 4, 'Note').font = B
def addname(n, ref): wb.defined_names[n] = DefinedName(n, attr_text=ref)
for i, (n, lab, f, note) in enumerate(rows):
    r = r0 + i; st.cell(r, 2, lab); st.cell(r, 3, f); st.cell(r, 4, note); addname(n, "'Plan stress test'!$C$%d" % r)
st['C4'].font = Font(color='0000FF', bold=True); st['C4'].fill = PatternFill('solid', fgColor='FFF2CC'); st['C4'].number_format = '0%'
GREY = Font(color='808080', size=9)
st['F3'] = 'Helper: age if short, after the fall'; st['G3'] = 'Helper: age if short, in your plan'; st['F3'].font = GREY; st['G3'].font = GREY
for i in range(90):
    st.cell(4 + i, 6, '=IF(\'Stress cashflow\'!$TV$%d=1,\'Stress cashflow\'!$D$%d,"")' % (8 + i, 8 + i)).font = GREY
    st.cell(4 + i, 7, '=IF(\'Plan cashflow\'!$TV$%d=1,\'Plan cashflow\'!$D$%d,"")' % (8 + i, 8 + i)).font = GREY
st.column_dimensions['F'].width = 14; st.column_dimensions['G'].width = 14
r = r0 + len(rows)
st.cell(r, 2, 'Summary sentence'); st.cell(r, 3, '="If markets fell "&ROUND(ST_Fall*100,0)&"% in "&ST_Year&" (age "&ST_Age&"): "&IF(ST_RunOutAge="","your money would still last to the end of your plan","your money would run short at "&ST_RunOutAge&IF(ST_RunOutAgeBase=""," (in your plan: it lasts to the end)",IF(ST_RunOutAgeBase=ST_RunOutAge,", the same as in your plan"," (in your plan: "&ST_RunOutAgeBase&")")))&"."')
addname('ST_Line', "'Plan stress test'!$C$%d" % r); st.cell(r, 3).alignment = Alignment(wrap_text=False)
r += 2
st.cell(r, 2, 'EACH GOAL (same order as Plan goals)').font = B; st.cell(r, 3, 'Covered after the fall').font = B; st.cell(r, 4, 'Line shown to the customer').font = B
for i in range(1, 23):
    r += 1
    st.cell(r, 2, f'=IF(\'Plan goals\'!$Q{109+i}="","",\'Plan goals\'!$C{109+i})')
    st.cell(r, 3, f"=IF('Stress goals'!$G${109+i}=\"\",\"\",'Stress goals'!$G${109+i})"); addname(f'ST{i}_Pct', "'Plan stress test'!$C$%d" % r)
    st.cell(r, 4, f'=IF(C{r}="","",IF(C{r}<GR{i}_Pct,"Falls from "&GR{i}_Pct&"% to "&C{r}&"% covered","Still "&C{r}&"% covered"))'); addname(f'ST{i}_Line', "'Plan stress test'!$D$%d" % r)
rEnd = r
r += 2
st.cell(r, 2, 'Live check: stress-year pension pot rebuilt by hand').font = B
st.cell(r + 1, 2, 'Start pot x (1 - fall) + contributions, or (start pot - draw) x (1 - fall) when drawing')
st.cell(r + 1, 3, "=INDEX('Stress pay & tax'!$S$234:$S$323,ST_T+1)")   # placeholder, replaced below
chk = ("=ROUND(IF(INDEX('Stress pay & tax'!$G$234:$G$323,ST_T+1)=1,INDEX('Stress pay & tax'!$S$234:$S$323,ST_T+1)*(1-ST_Fall)+INDEX('Stress pay & tax'!$Q$234:$Q$323,ST_T+1),"
       "IF(INDEX('Stress pay & tax'!$C$234:$C$323,ST_T+1)<F_V,INDEX('Stress pay & tax'!$S$234:$S$323,ST_T+1)*(1-ST_Fall),(INDEX('Stress pay & tax'!$Z$234:$Z$323,ST_T+1)-INDEX('Stress pay & tax'!$AC$234:$AC$323,ST_T+1))*(1-ST_Fall))),2)")
st.cell(r + 1, 3, chk)
st.cell(r + 2, 2, 'Pension pot in the stress run, that year'); st.cell(r + 2, 3, "=ROUND(INDEX('Stress pay & tax'!$AD$234:$AD$323,ST_T+1),2)")
st.cell(r + 3, 2, 'Check'); st.cell(r + 3, 3, f'=IF(ABS(C{r+1}-C{r+2})<0.01,"OK","CHECK")'); addname('ST_Check', "'Plan stress test'!$C$%d" % (r + 3))
for row in st.iter_rows(min_row=4, max_row=r + 3):
    for c in row:
        if c.column == 4: c.alignment = Alignment(wrap_text=True, vertical='top')
# order: Plan stress test after Plan profile; copies at the very end
order = [s for s in wb.sheetnames if s not in list(CL.values()) + ['Plan stress test']]
order.insert(order.index('Plan profile') + 1, 'Plan stress test')
order += list(CL.values())
wb._sheets = [wb[s] for s in order]
# Contents, Test notes, Report spec
rm = wb['READ ME']
rm['A20'] = 'Plan stress test'; rm['B20'] = 'If markets fell 20% in the year before retirement: goals covered, run-out age, pension at retirement (a side run; the plan does not change).'
rm['A20'].hyperlink = Hyperlink(ref='A20', location="'Plan stress test'!A1", display='Plan stress test'); rm['A20'].font = Font(color='0563C1', underline='single')
rm['A21'] = 'Stress copies'; rm['B21'] = 'HIDDEN at the end: Stress pay & tax, Stress holdings, Stress goals, Stress cashflow. Copies of the plan engine with only the stress-year growth changed.'
tn = wb['Test notes']
tn['B12'] = 'Stress-year pension pot rebuilt by hand (Plan stress test)'; tn['C12'] = '=IF(IFERROR(ST_Check="OK",FALSE()),"OK","CHECK")'
tn['B13'] = 'All live checks'; tn['B13'].font = B; tn['C13'] = '=IF(COUNTIF(C5:C12,"OK")=8,"ALL OK","CHECK")'; tn['C13'].font = B
tn['B22'] = 'Stress test'; tn['D22'] = 'Compared with the prototype\'s stressInfo() for each customer in the plan comparison: goals covered after the fall, run-out age and summary sentence.'
rs = wb['Report spec']; n = rs.max_row + 1
rs.cell(n, 2, 'Cashflow page, "How each year is paid" caption'); rs.cell(n, 3, 'ST_Line'); rs.cell(n, 4, 'Read: the stress-test sentence is added to the end of the caption (Plan stress test sheet, ST_Line).')
wb.save(out)
