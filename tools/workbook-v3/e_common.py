import openpyxl, re, json, copy
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
SC='/tmp/claude-0/-home-user-lifegoals-prototype/bb23c080-deb2-506f-b951-1207e03dac85/scratchpad/wb3/'
YEL=PatternFill('solid',fgColor='FFF2CC'); HEAD=PatternFill('solid',fgColor='1F4E79'); GRN=PatternFill('solid',fgColor='E2EFDA'); GREY=PatternFill('solid',fgColor='EDEDED'); PLN=PatternFill('solid',fgColor='DDEBF7')
BLUE=Font(color='0000FF'); WH=Font(color='FFFFFF',bold=True); BOLD=Font(bold=True); GREYF=Font(color='7F7F7F')
def qs(n): return "'"+n+"'"
def name(wb,n,ref):
    if n in wb.defined_names: del wb.defined_names[n]
    wb.defined_names[n]=DefinedName(n,attr_text=ref)
def cellname(wb,n,ws,cell):
    col=re.match(r'[A-Z]+',cell).group(0); row=cell[len(col):]
    name(wb,n,f"{qs(ws.title)}!${col}${row}")
