from lib import *

# ---- "Plan inputs": one yellow block mirroring the prototype's customer state ----
# kind: n = number (blank = 0, like the prototype's n()), t = text choice (blank = ""), a = an assumption the customer chose (blank = not chosen, engine falls back to the standard)
ABOUT = [
 ('age', 'Your age', 'age', 'n', 38, 'About you. The prototype: S.about.age'),
 ('partner', 'Planning with a partner? (Yes/No)', '', 't', 'Yes', 'S.about.partner'),
 ('married', 'Married or civil partners? (Yes/No, blank = not said)', '', 't', '', 'S.about.married (Yes = joint tax assessment)'),
 ('deps', 'Children or others who rely on you', 'number', 'n', 2, 'S.about.deps'),
 ('retireAge', 'Retirement age (your choice, never filled in for you)', 'age', 'n', 63, 'S.retireAge'),
 ('pRet', "Partner's retirement age (your choice; needed when your partner has income)", 'age', 'n', 66, 'S.pRet'),
 ('infl', 'Inflation you choose (a rate, e.g. 0.02 = 2%)', 'rate', 'n', '=Set_infl', 'S.infl'),
 ('assumeSet', 'Growth set (Standard / Cautious)', '', 't', 'Standard', 'S.assume'),
 ('q6', 'Q6 answer: how much could you invest? (0 to 4, blank = not answered)', '0-4', 'n', 2, "S.ans['6']"),
 ('saveM', 'My monthly saving (today\'s money, blank = use the rule)', '€ a month', 'n', '', 'S.saveM'),
 ('saveUp', 'My saving: raise the rule (Yes) or cap it (No); blank = decided by the amount', '', 't', '', 'S.saveUp'),
 ('penExtra', 'Extra monthly pension contribution you chose (€ a month)', '€ a month', 'n', 0, 'S.penExtra'),
 ('credRent', 'Claim the Rent Tax Credit? (Yes/No)', '', 't', 'No', 'S.about.cred.rent'),
 ('credLone', 'Claim the Single Person Child Carer Credit? (Yes/No)', '', 't', 'No', 'S.about.cred.lone'),
 ('credCarer', 'Claim the Home Carer Credit? (Yes/No)', '', 't', 'No', 'S.about.cred.carer'),
]
WHATIF = [
 ('wiGoal', 'What if: put it towards goal number (1 to 22, as listed below; blank = none)', 'goal no.', 'n', '', 'S.wi.g'),
 ('wiM', 'What if: extra (or less) each month', '€ a month', 'n', 0, 'S.wi.m'),
 ('wiL', 'What if: one-off amount', '€', 'n', 0, 'S.wi.l'),
]
FIN = [
 ('work', 'Your work (Employed / Self-employed / Not working)', '', 't', 'Employed', 'S.fin.work'),
 ('income', 'Your gross yearly income', '€', 'n', 58000, 'S.fin.income'),
 ('pIncome', "Partner's gross yearly income", '€', 'n', 22000, 'S.fin.pIncome'),
 ('pAge', "Partner's age", 'age', 'n', 37, 'S.fin.pAge'),
 ('otherM', 'Other income a month (before tax)', '€ a month', 'n', 0, 'S.fin.otherM'),
 ('costsM', 'Monthly living costs (leave out mortgage and loan repayments)', '€ a month', 'n', 3600, 'S.fin.costsM'),
 ('oneOffY', 'Yearly one-off costs', '€', 'n', 1800, 'S.fin.oneOffY'),
 ('propValue', 'Other property: value (only used for the section status)', '€', 'n', '', 'S.fin.propValue'),
 ('rentM', 'Other property: rent received a month', '€ a month', 'n', '', 'S.fin.rentM'),
 ('home', 'Your home (Own outright / Own with mortgage / Rent / Live with family)', '', 't', 'Own with mortgage', 'S.fin.home'),
 ('homeValue', 'Home value', '€', 'n', 420000, 'S.fin.homeValue (only counted as a missing detail)'),
 ('cash', 'Cash savings', '€', 'n', 9000, 'S.fin.cash'),
 ('invest', 'Investments', '€', 'n', 6200, 'S.fin.invest'),
 ('mortBal', 'Mortgage left', '€', 'n', 203700, 'S.fin.mortBal'),
 ('mortRate', 'Mortgage rate, APR (e.g. 0.0385 = 3.85%; blank = your assumption)', 'rate', 'n', 0.0385, 'S.fin.mortRate (the app shows it as a %, here a rate)'),
 ('mortPayM', 'Mortgage monthly repayment', '€ a month', 'n', 1180, 'S.fin.mortPayM'),
 ('mortYears', 'Mortgage years left', 'years', 'n', 21, 'S.fin.mortYears'),
 ('cardBal', 'Credit cards: total owed (a list on "Your lists" wins)', '€', 'n', 400, 'S.fin.cardBal'),
 ('cardPayM', 'Credit cards: monthly repayment', '€ a month', 'n', 40, 'S.fin.cardPayM'),
 ('cardRate', 'Credit cards: APR (e.g. 0.22 = 22%; blank = your assumption)', 'rate', 'n', '', 'S.fin.cardRate (the app shows it as a %, here a rate)'),
 ('loanBal', 'Other loans: total owed', '€', 'n', 3500, 'S.fin.loanBal'),
 ('loanPayM', 'Other loans: monthly repayment', '€ a month', 'n', 100, 'S.fin.loanPayM'),
 ('loanRate', 'Other loans: APR (e.g. 0.09 = 9%; blank = your assumption)', 'rate', 'n', '', 'S.fin.loanRate (the app shows it as a %, here a rate)'),
 ('debt', 'Other debt (older single figure)', '€', 'n', 0, 'S.fin.debt'),
 ('debtPayM', 'Other debt monthly repayment', '€ a month', 'n', 0, 'S.fin.debtPayM'),
 ('pension', 'Pension value today (a list on "Your lists" wins)', '€', 'n', 72000, 'S.fin.pension'),
 ('pensionM', 'Paid into your pension each month, you + employer', '€ a month', 'n', 520, 'S.fin.pensionM'),
 ('pensionOwnM', 'Of that, paid by you each month (blank = your share assumption)', '€ a month', 'n', '', 'S.fin.pensionOwnM'),
 ('penChg', "Your statement's pension charges, a rate (e.g. 0.01 = 1%; blank = the typical charges are already in the growth rate)", 'rate', 'n', '', 'S.fin.penChg (the app shows it as a %, here a rate)'),
 ('ae', 'Auto-enrolled in My Future Fund? (Yes/No/Not sure; "No" turns it off)', '', 't', '', 'S.fin.ae'),
 ('sp', 'State Pension (Expect full / Partly / Not sure)', '', 't', 'Expect full', 'S.fin.sp'),
 ('spYears', 'PRSI years at 66 (blank = not known)', 'years', 'n', '', 'S.fin.spYears'),
 ('pSp', "Partner's State Pension (Own full / Own partial / Qualified adult increase / None / Not sure)", '', 't', '', 'S.fin.pSp'),
 ('life', 'Life cover (Yes/No/Not sure)', '', 't', 'Yes', 'S.fin.life'),
 ('ip', 'Income protection (Yes/No/Not sure)', '', 't', 'No', 'S.fin.ip'),
 ('ci', 'Serious illness cover (Yes/No/Not sure)', '', 't', 'Not sure', 'S.fin.ci'),
 ('workCover', 'Cover through work (Yes/No/Not sure)', '', 't', 'Yes', 'S.fin.workCover'),
 ('health', 'Health insurance (Yes/No/Not sure)', '', 't', 'Yes', 'S.fin.health'),
]
# assumptions the customer can choose: (key, label, unit, example, prototype key)
ASMS = [
 ('a_inv', 'Investments grow (a year)', 'rate', '=Set_inv', 'inv'), ('a_cash', 'Cash savings earn (a year)', 'rate', '=Set_cash', 'cash'),
 ('a_pen', 'Pension grows while you work (a year)', 'rate', '=Set_pen', 'pen'), ('a_penRet', 'Pension grows once retired (a year)', 'rate', '=Set_penRet', 'penRet'),
 ('a_wage', 'Pay rises (a year)', 'rate', '=Set_wage', 'wage'), ('a_bands', 'Tax bands in future (prices / flat)', '', '=Set_bands', 'bands'),
 ('a_spGrow', 'State Pension in future (prices / pay / flat)', '', '=Set_spGrow', 'spGrow'), ('a_ownShare', 'Share of your pension you pay yourself', 'rate', '=IF(In_work="Self-employed",1,Set_ownShare)', 'ownShare'),
 ('a_accessAge', 'Pension access age (60 is usual; 50 only if your occupational scheme allows it)', 'age', '=Set_accessAge', 'accessAge'),
 ('a_retireSpend', 'Retirement spending if no retirement goal (share of today\'s)', 'rate', '=Set_retireSpend', 'retireSpend'), ('a_lumpSum', 'Tax-free lump sum to take (share of pension)', 'rate', '=Set_lumpSum', 'lumpSum'),
 ('a_drawRule', 'How you draw your pension (spread / min / fixed)', '', '=Set_drawRule', 'drawRule'), ('a_drawFixed', 'Fixed amount a year, if you chose fixed (today\'s money)', '€', '', 'drawFixed'),
 ('a_safetyMonths', 'Emergency fund months', 'months', '=IF(AS_TwoSecure=1,Set_safetyMonthsTwo,Set_safetyMonths)', 'safetyMonths'), ('a_cashYears', 'Money for goals sooner than this is kept as cash (years)', 'years', '=Set_cashYears', 'cashYears'),
 ('a_investShare', 'Spare savings: share invested', 'rate', '=Set_investShare', 'investShare'), ('a_saveShare', 'Share of spare money saved, at most', 'rate', '=Set_saveShare', 'saveShare'),
 ('a_noAnswer', 'Monthly saving if Q6 not answered', '€ a month', '=Set_noAnswer', 'noAnswer'), ('a_planEnd', 'Plan until age (your choice, never filled in for you)', 'age', 95, 'planEnd'),
 ('a_startYear', 'Start the plan in (2026 / 2027)', 'year', '=IF(TODAY()>PR_StartSwitch,Rules_Year_Next,Rules_Year)', 'startYear'),
 ('a_mortRate', 'Mortgage rate if the statement does not say', 'rate', 0.0375, 'mortRate'), ('a_cardRate', 'Credit-card rate if the statement does not say', 'rate', 0.2, 'cardRate'),
 ('a_loanRate', 'Other loans rate if the statement does not say', 'rate', 0.08, 'loanRate'), ('a_spWeek', 'Your State Pension a week at 66 (if you do not know your PRSI years)', '€ a week', '', 'spWeek'),
 ('a_pSpWeek', "Partner's State Pension a week at 66", '€ a week', 239.44, 'pSpWeek'), ('a_spYears', 'PRSI years at 66 (as an assumption)', 'years', '', 'spYears'),
 ('a_penChg', 'Typical pension charges already in the growth rate', 'rate', '', 'penChg'),
]
SESS = [  # session facts the app keeps that change what Home shows (not part of the plan numbers)
 ('booked', 'Adviser meeting booked? (Yes/No)', '', 't', '', 'S.book.done'),
 ('slot', 'Meeting slot as shown to you (text, e.g. Tue 14 Oct, 10:00)', '', 't', '', 'S.book.slot'),
 ('rechecked', 'Re-checked your first answers? (Yes/No)', '', 't', '', 'S.um.rechecked'),
 ('prefSet', 'Chose a what-if as your plan? (Yes/No)', '', 't', '', 'S.pref (set)'),
 ('saved1', 'Section saved: About you (Yes/No)', '', 't', '', 'S.saved.about'), ('saved2', 'Section saved: Income and expenses (Yes/No)', '', 't', '', 'S.saved.income'),
 ('saved3', 'Section saved: Assets (Yes/No)', '', 't', '', 'S.saved.assets'), ('saved4', 'Section saved: Liabilities (Yes/No)', '', 't', '', 'S.saved.liab'),
 ('saved5', 'Section saved: Protection (Yes/No)', '', 't', '', 'S.saved.prot'), ('saved6', 'Section saved: Pension (Yes/No)', '', 't', '', 'S.saved.pension'),
 ('lookKeys', 'Figures flagged "Needs a look" by a document or by you (field keys separated by commas, e.g. income,cash)', '', 't', 'pension', 'S.look keys (the pay-too-low flags are worked out for you)'),
 ('ackKeys', 'Flags you chose to "Leave for my adviser" (field keys separated by commas)', '', 't', '', 'S.ack keys'),
]
GOALKEYS = [  # GOALCAT: key, title, default amount, default years, kind
 ('home', 'Buy a home', 40000, 4, 'spend'), ('retire', 'Retire comfortably', 40000, 0, 'retire'), ('family', 'Grow my family', 15000, 3, 'spend'), ('edu', "Kids' education", 30000, 10, 'spend'),
 ('travel', 'Travel', 8000, 2, 'spend'), ('business', 'Start a business', 25000, 5, 'spend'), ('mfree', 'Be mortgage-free', 150000, 15, 'mfree'), ('safety', 'Emergency fund', 15000, 2, 'pot'),
 ('wealth', 'Grow my wealth', 50000, 10, 'pot'), ('helpfam', 'Help my family', 20000, 8, 'spend'), ('legacy', 'Leave a legacy', 100000, 0, 'legacy'), ('wedding', 'Wedding', 25000, 3, 'spend'),
 ('car', 'Change the car', 30000, 3, 'spend'), ('health', 'Health & care', 10000, 5, 'spend'), ('other', 'Something else', 10000, 4, 'spend')]
GOAL_EX = [  # the sample customer's goals, in order: key, age, amount, saved, prio, auto
 ('edu', 48, 50000, 0, 'Must have', 'No'), ('travel', 41, 12000, 0, 'Must have', 'No'), ('mfree', 55, 150000, 0, 'Must have', 'No'), ('safety', 40, 15000, 5000, 'Must have', 'No'), ('retire', 63, 46000, 0, 'Must have', 'No')]
NGOALS = 22   # goal rows: 20 funded goals + Retire + Leave a legacy

import json as _json
from openpyxl.worksheet.datavalidation import DataValidation
RJ = _json.load(open('/tmp/claude-0/-home-user-lifegoals-prototype/bb23c080-deb2-506f-b951-1207e03dac85/scratchpad/r15/ranges.json'))
YN = ['Yes', 'No']; YNN = ['Yes', 'No', 'Not sure']
EUR_MAX = 100000000   # a sanity ceiling for money figures; the app has no upper limit on these
# the app's own limits (FF fields, number-box limits and the ASM table): kind, low, high  |  list
SPEC = {
 'a_accessAge': ('list', [50, 60]),
 'age': ('whole', 18, 80), 'partner': ('list', YN), 'married': ('list', YN), 'deps': ('whole', 0, 20), 'retireAge': ('whole', 18, 105), 'pRet': ('whole', 18, 105),
 'infl': ('decimal', 0, 0.10), 'assumeSet': ('list', ['Standard', 'Cautious']), 'q6': ('whole', 0, 4), 'saveM': ('decimal', 0, 1000000), 'saveUp': ('list', YN), 'penExtra': ('decimal', 0, 100000),
 'credRent': ('list', YN), 'credLone': ('list', YN), 'credCarer': ('list', YN),
 'wiGoal': ('whole', 1, 22), 'wiM': ('decimal', -5000, 10000), 'wiL': ('decimal', 0, 2000000),
 'work': ('list', ['Employed', 'Self-employed', 'Not working']), 'pAge': ('whole', 18, 85), 'home': ('list', ['Own outright', 'Own with mortgage', 'Rent', 'Live with family']),
 'mortYears': ('whole', 1, 40), 'mortRate': ('decimal', 0, 0.30), 'cardRate': ('decimal', 0, 0.30), 'loanRate': ('decimal', 0, 0.30), 'penChg': ('decimal', 0, 0.05), 'spYears': ('whole', 0, 45),
 'life': ('list', YNN), 'ip': ('list', YNN), 'ci': ('list', YNN), 'workCover': ('list', YNN), 'health': ('list', YNN), 'ae': ('list', YNN),
 'sp': ('list', ['Expect full', 'Partly', 'Not sure']), 'pSp': ('list', ['Own full', 'Own partial', 'Qualified adult increase', 'None', 'Not sure']),
 'booked': ('list', YN), 'rechecked': ('list', YN), 'prefSet': ('list', YN),
 'saved1': ('list', YN), 'saved2': ('list', YN), 'saved3': ('list', YN), 'saved4': ('list', YN), 'saved5': ('list', YN), 'saved6': ('list', YN),
}
def spec_for(key, kind, unit, note):
    if key in SPEC: return SPEC[key]
    if kind == 'a':
        pk = note.split('.')[-1]; a = RJ['asm'].get(pk)
        if a is None: return None
        if a['t'] == 'choice': return ('list', a['opts'])
        sc = 100.0 if a['t'] == 'pct' else 1.0
        lo, hi = round(a['min'] / sc, 10), round(a['max'] / sc, 10)
        return ('whole' if a['t'] in ('yrs', 'age', 'num') and a.get('step', 1) % 1 == 0 else 'decimal', lo, hi)
    if unit.startswith('€'): return ('decimal', 0, EUR_MAX)
    return None
def valid_expr(spec, cell):
    if spec is None: return None
    if spec[0] == 'list':
        arr = ','.join(('"' + str(x) + '"') if isinstance(x, str) else str(x) for x in spec[1])
        if arr == '2026,2027': return f'OR({cell}=Rules_Year,{cell}=Rules_Year_Next)'   # the two plan start years are named cells on Assumptions
        return f'ISNUMBER(MATCH({cell},{{{arr}}},0))'
    lo, hi = spec[1], spec[2]
    e = f'IF(ISNUMBER({cell}),AND({cell}>={lo:g},{cell}<={hi:g}' + (f',{cell}=INT({cell})' if spec[0] == 'whole' else '') + '),FALSE)'   # text never reaches INT() (that gave an error value)
    return e
def add_dv(ws, spec, cell, label):
    if spec is None: return
    if spec[0] == 'list':
        lst = ','.join(str(x) for x in spec[1])
        dv = DataValidation(type='list', formula1='"' + lst + '"', allow_blank=True, showErrorMessage=True, errorTitle='Not in the list', error='Choose one of: ' + lst)
    else:
        dv = DataValidation(type=spec[0], operator='between', formula1=f'{spec[1]:g}', formula2=f'{spec[2]:g}', allow_blank=True, showErrorMessage=True, errorTitle='Outside the allowed range',
                            error=f'{label}: allowed {spec[1]:g} to {spec[2]:g}' + (' (whole numbers)' if spec[0] == 'whole' else ''))
    ws.add_data_validation(dv); dv.add(cell)

def build_inputs(bk):
    wb = bk.wb
    if 'Plan inputs' in wb.sheetnames: del wb['Plan inputs']
    ws = wb.create_sheet('Plan inputs'); ws.sheet_properties.tabColor = '5B9BD5'
    ws['B1'] = 'Plan inputs: everything the "Make my plan" engine needs (same as the prototype\'s customer state)'; ws['B1'].font = Font(bold=True, size=14)
    ws['B2'] = ('Type your figures in the yellow cells (column C). A blank cell is €0, or "not chosen" for a choice: the plan then uses the standard from Settings and "Plan results" lists what still needs choosing '
                '("we need N things"). Switch "Fill example values" on README to Yes and every blank cell uses the grey example (the prototype\'s sample customer, Aoife and Cian). '
                'Column G is the figure the plan uses; column H says 1 if you gave it. Never type in columns F, G or H.')
    ws['B2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells('B2:I2'); ws.row_dimensions[2].height = 62
    for i, w in enumerate([2, 66, 16, 12, 38, 16, 16, 8, 40]): ws.column_dimensions[CL(i + 1)].width = w
    r = 4
    def head(t):
        nonlocal r
        for j, h in enumerate([t, 'Your figure', 'Unit', 'Prototype state key', 'Example (sample)', 'Used by the plan', 'Given?']):
            c = ws.cell(r, 2 + j, h); c.fill = HEAD; c.font = WH
        r += 1
    VALIDS = {}
    def row(key, label, unit, kind, ex, note, usedf=None):
        nonlocal r
        ws.cell(r, 2, label); ws.cell(r, 4, 'rate: 0.0375 = 3.75%' if unit == 'rate' else unit); ws.cell(r, 5, note).font = GREYF
        c = ws.cell(r, 3); c.fill = YEL; c.font = BLUE
        exc = ws.cell(r, 6, ex if ex != '' else None); exc.font = GREYF
        sp = spec_for(key, kind, unit, note); V = valid_expr(sp, f'C{r}') or 'TRUE'
        if kind == 'n':
            g = f'=IF(ISBLANK(C{r}),IF(AND(Fill_Example="Yes",F{r}<>""),F{r},0),IF({V},C{r},0))'
            h = f'=IF(ISBLANK(C{r}),IF(AND(Fill_Example="Yes",F{r}<>""),1,0),IF({V},1,0))'
        elif kind == 't':
            g = f'=IF(ISBLANK(C{r}),IF(Fill_Example="Yes",F{r},""),IF({V},C{r},""))'
            h = f'=IF(G{r}="",0,1)'
        else:  # a
            g = f'=IF(ISBLANK(C{r}),IF(AND(Fill_Example="Yes",F{r}<>""),F{r},""),IF({V},C{r},""))'
            h = f'=IF(G{r}="",0,1)'
        add_dv(ws, sp, f'C{r}', label)
        if usedf: g = usedf.replace('{r}', str(r)).replace('{V}', V)
        VALIDS[key] = V
        ws.cell(r, 7, g).fill = GRN; ws.cell(r, 8, h).fill = GRN
        bk.cellname('In_' + key, ws, f'G{r}'); bk.cellname('Has_' + key, ws, f'H{r}')
        r += 1
    head('ABOUT YOU, CHOICES AND SAVING')
    for k in ABOUT: row(*k)
    r += 1; head('WHAT IF (Results screen)')
    for k in WHATIF: row(*k)
    r += 1; head('YOUR FINANCES')
    merged = {  # list totals win when any item has a figure (as syncLists does)
      'cardBal': '=IF(COUNT(Cards_Owed_Rng)>0,SUM(Cards_Owed_Rng),IF(ISBLANK(C{r}),IF(AND(Fill_Example="Yes",F{r}<>""),F{r},0),IF({V},C{r},0)))',
      'cardPayM': '=IF(COUNT(Cards_Pay_Rng)>0,SUM(Cards_Pay_Rng),IF(ISBLANK(C{r}),IF(AND(Fill_Example="Yes",F{r}<>""),F{r},0),IF({V},C{r},0)))',
      'cardRate': '=IF(ISNUMBER(Cards_Rate_Avg),ROUND(Cards_Rate_Avg,6),IF(ISBLANK(C{r}),IF(AND(Fill_Example="Yes",F{r}<>""),F{r},0),IF({V},C{r},0)))',
      'loanBal': '=IF(COUNT(Loans_Owed_Rng)>0,SUM(Loans_Owed_Rng),IF(ISBLANK(C{r}),IF(AND(Fill_Example="Yes",F{r}<>""),F{r},0),IF({V},C{r},0)))',
      'loanPayM': '=IF(COUNT(Loans_Pay_Rng)>0,SUM(Loans_Pay_Rng),IF(ISBLANK(C{r}),IF(AND(Fill_Example="Yes",F{r}<>""),F{r},0),IF({V},C{r},0)))',
      'loanRate': '=IF(ISNUMBER(Loans_Rate_Avg),ROUND(Loans_Rate_Avg,6),IF(ISBLANK(C{r}),IF(AND(Fill_Example="Yes",F{r}<>""),F{r},0),IF({V},C{r},0)))',
      'pension': '=IF(COUNT(Pens_Value_Rng)>0,SUM(Pens_Value_Rng),IF(ISBLANK(C{r}),IF(AND(Fill_Example="Yes",F{r}<>""),F{r},0),IF({V},C{r},0)))',
      'pensionM': '=IF(COUNT(Pens_Monthly_Rng)>0,SUM(Pens_Monthly_Rng),IF(ISBLANK(C{r}),IF(AND(Fill_Example="Yes",F{r}<>""),F{r},0),IF({V},C{r},0)))',
      'pensionOwnM': '=IF(COUNT(Pens_Monthly_Rng)>0,IF(Pens_OwnGiven=1,Pens_Own_Total,0),IF(ISBLANK(C{r}),IF(AND(Fill_Example="Yes",F{r}<>""),F{r},0),IF({V},C{r},0)))',
    }
    for k in FIN:
        usedf = merged.get(k[0]); row(*k, usedf=usedf)
    # Has_ for list-merged keys: a list entry counts as given
    for k, rng in [('cardBal', 'Cards_Owed_Rng'), ('cardPayM', 'Cards_Pay_Rng'), ('loanBal', 'Loans_Owed_Rng'), ('loanPayM', 'Loans_Pay_Rng'), ('pension', 'Pens_Value_Rng'), ('pensionM', 'Pens_Monthly_Rng')]:
        rr = int(bk.wb.defined_names['Has_' + k].attr_text.split('$')[-1]); ws.cell(rr, 8).value = f'=IF(COUNT({rng})>0,1,IF(ISBLANK(C{rr}),IF(AND(Fill_Example="Yes",F{rr}<>""),1,0),IF({VALIDS[k]},1,0)))'
    rr = int(bk.wb.defined_names['Has_pensionOwnM'].attr_text.split('$')[-1])
    ws.cell(rr, 8).value = f'=IF(COUNT(Pens_Monthly_Rng)>0,Pens_OwnGiven,IF(ISBLANK(C{rr}),IF(AND(Fill_Example="Yes",F{rr}<>""),1,0),IF({VALIDS["pensionOwnM"]},1,0)))'
    r += 1; head('SESSION FACTS (what Home and My money show; none of them changes the plan numbers)')
    for k in SESS: row(*k)
    r += 1; head('YOUR ASSUMPTIONS (the choices the customer makes; blank = not chosen yet)')
    for key, label, unit, ex, pk in ASMS: row(key, label, unit, 'a', ex, 'S.asm.' + pk)
    # goals table
    r += 1
    gh = ['YOUR GOALS (up to 22: 20 saved for from savings, plus Retire and Leave a legacy; in the order you added them)', 'Goal type (key)', 'Name (blank = standard name)', 'Target age (not used for Retire or Legacy)', 'Amount today (Retire: yearly income)', 'Saved so far', 'Priority', 'Rank (1 = first; blank = not ranked)', 'Emergency fund amount follows months x spending? (Yes/No)']
    for j, h in enumerate(gh):
        c = ws.cell(r, 2 + j, h); c.fill = HEAD; c.font = WH; c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[r].height = 62
    ws.column_dimensions['J'].width = 14; ws.column_dimensions['K'].width = 14; ws.column_dimensions['L'].width = 14
    r += 1; g0 = r
    # used block to the right (columns N..U) holds the example when the key cell is blank and Fill example = Yes
    ws.cell(g0 - 1, 14, 'Used: key').fill = HEAD
    for j, h in enumerate(['Used: name', 'Used: age', 'Used: amount', 'Used: saved', 'Used: priority', 'Used: rank', 'Used: auto']): ws.cell(g0 - 1, 15 + j, h)
    for j in range(8): ws.cell(g0 - 1, 14 + j).fill = HEAD; ws.cell(g0 - 1, 14 + j).font = WH
    for i in range(NGOALS):
        rr = g0 + i; ex = GOAL_EX[i] if i < len(GOAL_EX) else None
        ws.cell(rr, 2, f'Goal {i + 1}')
        for j in range(8):
            c = ws.cell(rr, 3 + j); c.fill = YEL; c.font = BLUE
        # typed in C..J (key, name, age, amount, saved, prio, rank, auto); used in N..U
        exk = ex[0] if ex else ''
        kk = f'IF(ISBLANK(C{rr}),IF(Fill_Example="Yes","{exk}",""),IF(ISNUMBER(MATCH(C{rr},Goal_Keys,0)),C{rr},""))'
        add_dv(ws, ('list', [g[0] for g in GOALKEYS]), f'C{rr}', 'Goal type')
        GS = {'E': ('whole', 0, 105), 'F': ('decimal', 0, EUR_MAX), 'G': ('decimal', 0, EUR_MAX), 'H': ('list', ['Must have', 'Nice to have']), 'I': ('whole', 1, NGOALS), 'J': ('list', YN)}
        for col, spc in GS.items(): add_dv(ws, spc, f'{col}{rr}', {'E': 'Target age', 'F': 'Amount', 'G': 'Saved so far', 'H': 'Priority', 'I': 'Rank', 'J': 'Emergency fund follows spending'}[col])
        ws.cell(rr, 14, '=' + kk)
        isEx = f'AND(ISBLANK(C{rr}),Fill_Example="Yes")'
        def used(col, exv, default):
            V = valid_expr(GS[col], f'{col}{rr}')
            return f'=IF(ISBLANK({col}{rr}),IF({isEx},{exv},{default}),IF({V},{col}{rr},{default}))'
        exv = lambda x: ('"' + x + '"') if isinstance(x, str) else str(x)
        t = lambda col: f'N{rr}'
        title = f'IFERROR(INDEX(Goal_Titles,MATCH(N{rr},Goal_Keys,0)),"")'
        defamt = f'IFERROR(INDEX(Goal_DefAmt,MATCH(N{rr},Goal_Keys,0)),0)'
        defage = f'MIN(PR_GoalMaxAge,In_age+IFERROR(INDEX(Goal_DefYrs,MATCH(N{rr},Goal_Keys,0)),0))'
        ws.cell(rr, 15, f'=IF(ISBLANK(D{rr}),{title},D{rr})')
        ws.cell(rr, 16, used('E', exv(ex[1]) if ex else 0, defage))
        ws.cell(rr, 17, used('F', exv(ex[2]) if ex else 0, defamt))
        ws.cell(rr, 18, used('G', exv(ex[3]) if ex else 0, 0))
        ws.cell(rr, 19, used('H', exv(ex[4]) if ex else '"Must have"', '"Must have"'))
        ws.cell(rr, 20, f'=IF(ISBLANK(I{rr}),"",IF({valid_expr(GS["I"], f"I{rr}")},I{rr},""))')
        ws.cell(rr, 21, used('J', exv(ex[5]) if ex else '"No"', '"No"'))
        for j, nm in enumerate(['key', 'name', 'age', 'amount', 'saved', 'prio', 'rank', 'auto']):
            bk.cellname(f'Goal{i + 1}_{nm}', ws, f'{CL(14 + j)}{rr}')
        for j in range(8): ws.cell(rr, 14 + j).fill = GRN
    bk.name('Goal_Rows_First', f"{qs('Plan inputs')}!$N${g0}")
    from openpyxl.formatting.rule import FormulaRule
    ws['B3'] = '=Warn_Goals'; ws['B3'].font = Font(bold=True, color='9C0006'); ws.merge_cells('B3:I3'); ws['B3'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.conditional_formatting.add('B3:I3', FormulaRule(formula=['LEN($B$3)>0'], fill=PatternFill('solid', bgColor='FFC7CE', fgColor='FFC7CE'), font=Font(bold=True, color='9C0006')))
    ws.row_dimensions[3].height = 30
    ws.freeze_panes = 'C4'
    return ws, g0
