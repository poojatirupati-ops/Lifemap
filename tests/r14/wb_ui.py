import openpyxl
P = '/home/user/lifegoals-prototype/deliverables/LifeGoals-Calculators.xlsx'
wb = openpyxl.load_workbook(P)
def dv_of(ws, cell):
    for dv in ws.data_validations.dataValidation:
        if str(dv.sqref) == cell: return dv
    raise SystemExit('no DV ' + cell)
def fixtxt(ws, cell, a, b):
    v = ws[cell].value; assert a in str(v), (ws.title, cell, v); ws[cell] = v.replace(a, b)
# (1) ranges = the app
ws = wb['C10 Lump sum growth']; dv = dv_of(ws, 'C20'); dv.formula2 = '20'; dv.error = 'Please enter 1 to 20 · steps of 1.'; fixtxt(ws, 'E20', 'Allowed: 1 to 10', 'Allowed: 1 to 20')
ws = wb['C11 Emergency fund']; dv = dv_of(ws, 'C7'); dv.formula1 = '1'; dv.error = 'Please enter 1 to 12 · steps of 1.'
v = ws['E7'].value; assert 'Allowed: 3 to 12' in str(v), v; ws['E7'] = v.replace('Allowed: 3 to 12', 'Allowed: 1 to 12')
# (2) steps = the app
for s in ('C13 Contribution impact', 'C14 AVC impact'):
    ws = wb[s]; dv_of(ws, 'C18').error = 'Please enter 1% to 7% · steps of 0.25%.'; fixtxt(ws, 'E18', 'steps of 0.5%', 'steps of 0.25%')
ws = wb['C19 Fees impact']
for c in ('C15', 'C16'): dv_of(ws, c).error = 'Please enter 0% to 3% · steps of 0.05%.'; fixtxt(ws, 'E' + c[1:], 'steps of 0.1%', 'steps of 0.05%')
ws = wb['C28 Loan repayment']; dv_of(ws, 'C7').error = 'Please enter 0% to 25% · steps of 0.25%.'; fixtxt(ws, 'E7', 'steps of 0.5%', 'steps of 0.25%')
# (3) take-home pay: the app starts the tool at an example figure of €3,500 (journey spec §26: example figures are labelled, never the customer's); the workbook keeps it blank until "Fill example values" = Yes
for s in ('C25 Monthly surplus', 'C26 Budget 50 30 20'):
    ws = wb[s]; assert ws['G6'].value == '=IF(ISBLANK(C6),"",C6)'
    ws['F6'] = 3500; ws['G6'] = '=IF(ISBLANK(C6),IF(Fill_Example="Yes",F6,""),C6)'
# (5) every grey example column is labelled
n = 0
for ws in wb:
    if not ws.title.startswith('C') or ws.title[1:3].isdigit() is False: continue
    for r in range(1, 30):
        if ws.cell(r, 3).value == 'Your figure' and ws.cell(r, 6).value in (None, ''):
            ws.cell(r, 6).value = 'Example only (grey; used only when "Fill example values" is Yes)'; ws.cell(r, 6).font = openpyxl.styles.Font(color='7F7F7F', italic=True); n += 1
print('example headers', n)
wb.save(P); print('ok')
