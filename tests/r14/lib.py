import re, json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as CL

YEL = PatternFill('solid', fgColor='FFF2CC'); HEAD = PatternFill('solid', fgColor='1F4E79'); GREY = PatternFill('solid', fgColor='EDEDED')
GRN = PatternFill('solid', fgColor='E2EFDA'); PLN = PatternFill('solid', fgColor='DDEBF7'); ORG = PatternFill('solid', fgColor='FCE4D6')
BLUE = Font(color='0000FF'); WH = Font(color='FFFFFF', bold=True); B = Font(bold=True); GREYF = Font(color='7F7F7F')
RI = json.load(open('/tmp/claude-0/-home-user-lifegoals-prototype/bb23c080-deb2-506f-b951-1207e03dac85/scratchpad/r14/ri.json'))

def qs(name): return "'" + name + "'"

class Book:
    def __init__(self, wb): self.wb = wb
    def name(self, n, ref):
        if n in self.wb.defined_names: del self.wb.defined_names[n]
        self.wb.defined_names[n] = DefinedName(n, attr_text=ref)
    def cellname(self, n, ws, cell):
        col = re.match(r'[A-Z]+', cell).group(0); row = cell[len(col):]
        self.name(n, f"{qs(ws.title)}!${col}${row}")

class YT:
    """A year table: one row per plan year t = 0..NROWS-1, columns defined by key with a formula template.
    {key} = same row, {key@-1} = previous row (the init row for t=0), {Alias.key} = another year table, same row."""
    reg = {}
    def __init__(self, ws, alias, hdr_row, nrows):
        self.ws, self.alias, self.hdr, self.n = ws, alias, hdr_row, nrows
        self.init_row = hdr_row + 1; self.r0 = hdr_row + 2; self.cols = []; self.idx = {}
        YT.reg[alias] = self
    def add(self, key, header, formula, init=None, fmt=None, width=11, note=None):
        assert key not in self.idx, key
        self.idx[key] = len(self.cols) + 2; self.cols.append(dict(key=key, header=header, f=formula, init=init, fmt=fmt, width=width, note=note))
    def col(self, key): return CL(self.idx[key])
    def ref(self, key, off, row):
        c = self.col(key); return f"{c}{row + off}"
    def resolve(self, text, row, own=None):
        def sub(m):
            tgt, off = m.group(1), m.group(2)
            if '.' in tgt: al, key = tgt.split('.', 1); yt = YT.reg[al]
            else: yt, key = self, tgt
            tidx = row - self.r0
            r = yt.ref(key, -1 if off else 0, yt.r0 + tidx)
            return (f"{qs(yt.ws.title)}!" if yt is not self else '') + r
        return re.sub(r'\{([A-Za-z0-9_.]+)(@-1)?\}', sub, text)
    def write(self):
        ws = self.ws
        ws.cell(self.hdr, 1, 'row')
        for c in self.cols:
            x = ws.cell(self.hdr, self.idx[c['key']], c['header']); x.fill = HEAD; x.font = WH; x.alignment = Alignment(wrap_text=True, vertical='top')
            ws.column_dimensions[CL(self.idx[c['key']])].width = c['width']
        ws.cell(self.hdr, 1).fill = HEAD; ws.cell(self.hdr, 1).font = WH
        ws.cell(self.init_row, 1, 'start')
        for c in self.cols:
            if c['init'] is not None:
                v = c['init']; v = self.resolve(v, self.init_row) if isinstance(v, str) else v
                x = ws.cell(self.init_row, self.idx[c['key']], v); x.fill = GREY
                if c['fmt']: x.number_format = c['fmt']
        for t in range(self.n):
            r = self.r0 + t
            for c in self.cols:
                f = c['f']
                if c['key'] == 't': v = t
                else: v = self.resolve(f, r)
                x = ws.cell(r, self.idx[c['key']], v)
                if c['fmt']: x.number_format = c['fmt']
        ws.freeze_panes = ws.cell(self.r0, 3)
