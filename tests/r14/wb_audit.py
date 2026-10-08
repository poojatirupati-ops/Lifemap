import openpyxl, re
from openpyxl.workbook.defined_name import DefinedName
P = '/home/user/lifegoals-prototype/deliverables/LifeGoals-Calculators.xlsx'
wb = openpyxl.load_workbook(P)
def setf(ws, addr, f):
    ws[addr].value = f
# ---- C12
ws = wb['C12 Retirement projection']
def chk(addr, start):
    v = ws[addr].value; assert str(v).startswith(start), (addr, v)
chk('C48', '=IF(Need_4'); ws['C48'] = '=IF(Need_4<>"",Need_4,MAX(0,Desired_Income-Spend_Less-Other_Income)*Retire_Multiple+Bridge_Years*MIN(Other_Income,MAX(0,Desired_Income-Spend_Less)))'
ws['E48'] = 'ƒx  =MAX(0,Desired_Income-Spend_Less-Other_Income)*Retire_Multiple+Bridge_Years*MIN(Other_Income,MAX(0,Desired_Income-Spend_Less))      (the years before the State Pension never need more than the income you need)'
chk('C50', '=IF(Need_5'); ws['C50'] = '=IF(Need_5<>"",Need_5,Projected_Fund*MIN(Lump_Sum_Pct,Lump_Sum_Max_Pct))'
ws['E50'] = 'ƒx  =Projected_Fund*MIN(Lump_Sum_Pct,Lump_Sum_Max_Pct)      (25% at most of the fund: no cap; first €200,000 tax-free, next €300,000 at 20%, the excess at the higher rate)'
chk('C51', '=IF(Need_5'); ws['C51'] = '=IF(Need_5<>"",Need_5,Lump_Sum_Rate*MAX(0,MIN(Lump_Sum,Lump_Sum_Cap)-Lump_Sum_Tax_Free)+Rate_Higher*MAX(0,Lump_Sum-Lump_Sum_Cap))'
ws['E51'] = 'ƒx  =Lump_Sum_Rate*MAX(0,MIN(Lump_Sum,Lump_Sum_Cap)-Lump_Sum_Tax_Free)+Rate_Higher*MAX(0,Lump_Sum-Lump_Sum_Cap)      (the part above €500,000 is taxed at the higher rate; verify on Revenue)'
chk('C54', '=IF(Need_3'); ws['C54'] = '=IF(Need_3<>"","",IF(AND(ISNUMBER(Fund_Today),IF(ISNUMBER(Fund_Today),Fund_Today,0)>SFT_2026),"Your fund in today\'s money is above the Standard Fund Threshold (€2.2m in 2026, rising to €2.8m by 2029; later years are not set yet): the excess is taxed at "&TEXT(SFT_Tax,"0%")&".",""))'
ws['B63'] = '• Desired income and the State Pension are before tax: the State Pension is taxable (USC-exempt) and drawdown pays income tax and USC. The lump sum is the share you chose (25% at most): the first €200,000 tax-free, the next €300,000 at 20%, anything above €500,000 at the higher rate. Contributions are counted as paid at the end of each year.'
ws['B65'] = '• A retirement age at or before your age gives "Choose a retirement age after your age to see this" (never a silent 1 year).'
ws['B62'] = "• Target = (desired income − spend less − other income) × your multiple (suggested 25), plus the State Pension you must fund yourself for each year you retire before 66 (at most the income you need each year). Gap = target − fund in today's money."
for r in range(77, 86):
    v = ws.cell(r, 3).value; assert str(v).startswith('=IF('), (r, v)
    ws.cell(r, 3).value = '=IF(AND(ISNUMBER(Age),ISNUMBER(Retirement_Age),N(Retirement_Age)+N(Retire_Later)<=N(Age)),"Choose a retirement age after your age to see this",' + v[1:] + ')'
# ---- C22
ws = wb['C22 Income protection gap']
assert 'Indefinitely' in ws['C13'].value
ws['C13'] = '=IF(Need_1<>"",Need_1,IF(Monthly_Gap=0,24,MIN(24,Sick_Pay_Months+Savings/Monthly_Gap)))'
ws['D13'] = '=IF(Need_1<>"",Need_1,IF(Monthly_Gap=0,"up to 2 years",IF(Sick_Pay_Months+Savings/Monthly_Gap>=24,"up to 2 years",IF(ROUND(Months_Coping,1)=1,"1 month",FIXED(Months_Coping,1)&" months"))))'
ws['B13'] = 'You could cope for (months, at most 24)'
ws['C17'] = '=IF(Need_5<>"",Need_5,IF(Monthly_Gap=0,"Not needed",Savings/Monthly_Gap))'
ws['D17'] = '=IF(Need_5<>"",Need_5,IF(Savings_Months="Not needed","Not needed while Illness Benefit covers your spending",IF(Savings_Months>=24,"more than 2 years",IF(ROUND(Savings_Months,1)=1,"1 month",FIXED(Savings_Months,1)&" months"))))'
ws['E13'] = 'ƒx  =IF(Monthly_Gap=0,24,MIN(24,Sick_Pay_Months+Savings/Monthly_Gap))      (Illness Benefit is paid for at most 2 years, 624 days)'
ws['B22'] = '• Months you could cope = sick-pay months + savings ÷ monthly gap, never more than 24 (Illness Benefit is paid for at most 2 years, 624 days).'
ws['B25'] = '• If Illness Benefit covers the spending, the time is shown as "up to 2 years" and savings as "Not needed" (never "Indefinitely").'
# ---- C15: less than a year, and the early-access message
ws = wb['C15 Will my money last']
f = ws['C29'].value; assert 'Beyond' in f
ws['C29'] = '=IF(Need_2<>"",Need_2,IF(ISNUMBER(Years_Lasting),IF(Years_Lasting>=Cap_Years,"Beyond "&Plan_To_Age,IF(Years_Lasting<1,"Less than a year",Age_At_Start+Years_Lasting)),"Choose your inflation rate to see this"))'
d = ws['D29'].value
assert '&" years before your plan age")),' in d
d = d.replace('(your plan-until age)",IF(Age_At_Start+Years_Lasting>=Plan_To_Age,', '(your plan-until age)",IF(Years_Lasting<1,"Less than a year after you start",IF(Age_At_Start+Years_Lasting>=Plan_To_Age,').replace('&" years before your plan age")),', '&" years before your plan age"))),')
ws['D29'] = d
d28 = ws['D28'].value; ws['D28'] = d28.replace('Years_Lasting&IF(Years_Lasting=1," year"," years")', 'IF(Years_Lasting<1,"Less than a year",Years_Lasting&IF(Years_Lasting=1," year"," years"))')
# ---- C07 deposit clamp
ws = wb['C07 Rent vs buy']
ws['B18'] = 'Deposit used (at least 10% of the price, at most the price)'
ws['C18'] = '=IF(OR(Home_Price="",Deposit=""),"",MIN(MAX(Deposit,Home_Price*(1-LTV_Max)),Home_Price))'
ws['D18'] = '€'
ws.defined_names['Deposit_Used'] = DefinedName('Deposit_Used', attr_text="'C07 Rent vs buy'!$C$18")
for a, old, new in [('C24', 'Home_Price-Deposit', 'Home_Price-Deposit_Used'), ('C30', 'Deposit*((1+Opp_Rate)', 'Deposit_Used*((1+Opp_Rate)'), ('C42', 'Deposit+Stamp_Duty', 'Deposit_Used+Stamp_Duty')]:
    v = ws[a].value; assert old in v, (a, v); ws[a] = v.replace(old, new)
# ---- notes
def add_note(sheet, text):
    w = wb[sheet]; r = w.max_row + 2; w.cell(r, 2, text); w.cell(r, 2).font = openpyxl.styles.Font(color='7F7F7F')
add_note('C01 Borrowing', 'Note: stamp duty is worked out on the price shown; for a new home it is charged on the price excluding VAT, and Help to Buy is not included.')
add_note('C03 Deposit', 'Note: stamp duty is worked out on the price shown; for a new home it is charged on the price excluding VAT, and Help to Buy is not included.')
add_note('C07 Rent vs buy', 'Note: stamp duty is worked out on the price shown; for a new home it is charged on the price excluding VAT, and Help to Buy is not included. Home equity is before selling costs.')
add_note('C02 Mortgage repayment', 'Note: the rate is a yearly rate, compounded monthly (nominal rate ÷ 12), the way lenders quote mortgage rates.')
add_note('C28 Loan repayment', 'Note: the APR is treated as an effective yearly rate (as for cards and loans); mortgage tools use the nominal rate ÷ 12, so the same percentage gives a slightly different repayment.')
add_note('C13 Contribution impact', 'Note: the extra contribution is the same every year; it does not rise with pay (the Retirement projection and your plan do let contributions rise with pay).')
add_note('C10 Lump sum growth', 'Note: the money is assumed to earn nothing while it waits. If growth after fees is zero or negative there is no cost of waiting to show.')
add_note('C20 Risk and return', 'Note: over longer timeframes the yearly return range narrows, but the gap in euros between the weaker and stronger outcomes widens.')
wa = wb['Assumptions']; wa['B55'] = 'Lump sum: taxed at 20% from €200,001 up to  [Lump_Sum_Cap]; above this the excess is taxed at the higher rate [Rate_Higher] (verify on Revenue)'
wb.save(P); print('ok')
