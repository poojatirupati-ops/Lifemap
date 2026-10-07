import openpyxl
P='/home/user/lifegoals-prototype/deliverables/LifeGoals-Calculators.xlsx'
wb=openpyxl.load_workbook(P);ws=wb['Settings']
r=int(wb.defined_names['Set_inv'].attr_text.split('$')[-1])
chg='IF(ISNUMBER(Suggest_Fund_Charges),Suggest_Fund_Charges,0.01)'
# same single wording as the app (invBy): built from the same inputs
ws.cell(r,7).value=('="a planning assumption, not a prediction: "&ROUND(Set_invGross*1000,0)/10&"% before fees and tax, less about "&ROUND('+chg+'*1000,0)/10&"% charges, less "&ROUND(Exit_Tax*100,0)&"% exit tax taken every "&Deemed_Disposal_Years&" years"')
ws.cell(r,10).value='A planning assumption, not a prediction. Growth is never guaranteed and funds can fall in value. It moves if you change the growth before fees and tax, or the charges. (Edit Set_invGross and the fund charges, not this cell.)'
wb.save(P);print('ok')
